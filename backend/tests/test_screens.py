"""Endpoints added for Screens 2-5: customers, channels, Excel exports, scenarios, and the Capital Adequacy report
(calendar, open report, validation, drill-to-source, PDF and Excel export)."""
import io
import re
from decimal import Decimal

import psycopg2
import psycopg2.errors
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.reports.form import ALLOCATION, FORM, FORM_CODES, allocate
from app.reports.formatting import line_text, money, percent
from app.reports.validation import CHECKS, run_validation

API = "/api/v1"


# ---- helpers --------------------------------------------------------------------------------
@pytest.fixture(scope="module")
def ca_id(client, auth):
    """The current-quarter Capital Adequacy instance (the only one with figures)."""
    rows = client.get(f"{API}/reports", headers=auth).json()
    return next(r["report_instance_id"] for r in rows if r["name"] == "Capital Adequacy")


@pytest.fixture(scope="module")
def report(client, auth, ca_id):
    return client.get(f"{API}/reports/{ca_id}", headers=auth).json()


def values(report):
    return {l["line_code"]: Decimal(str(l["value"])) for s in report["sections"] for l in s["lines"]}


# ---- pure logic -----------------------------------------------------------------------------
def test_money_and_percent_formatting():
    assert money(150000000) == "150,000,000"
    assert money(Decimal("-8000000.0000")) == "(8,000,000)"
    assert money(Decimal("0.5")) == "1" and money(Decimal("-0.4")) == "0"
    assert percent(Decimal("10.4282")) == "10.43%" and percent(12) == "12.00%"
    assert line_text("percent", Decimal("12.3929")) == "12.39%" and line_text("currency", 7800000) == "7,800,000"


@pytest.mark.parametrize("total", [207_000_000, 1_000_001, 999_999_999, 1])
def test_allocation_always_adds_back_to_the_real_total(total):
    for total_code, weights in ALLOCATION.items():
        assert sum(allocate(Decimal(total), weights).values()) == total


def test_every_form_line_has_a_check_that_can_run_and_the_check_registry_is_complete():
    assert len(FORM) == 17 and FORM_CODES[0] == "A.1" and FORM_CODES[-1] == "C.4"
    assert set(CHECKS) == {"TIER_TOTAL", "REQUIRED_FIELDS", "SECTION_TOTALS", "CAPITAL_BUFFER"}


# ---- report calendar ------------------------------------------------------------------------
def test_calendar_lists_example_rows_soonest_first_and_only_capital_adequacy_has_figures(client, auth):
    rows = client.get(f"{API}/reports", headers=auth).json()
    assert len(rows) == 8
    dues = [r["due_date"] for r in rows]
    assert dues == sorted(dues)
    assert [r["name"] for r in rows if r["has_lines"]] == ["Capital Adequacy"]
    assert {r["status"] for r in rows} == {"NOT_STARTED", "DRAFT", "UNDER_REVIEW", "APPROVED", "SUBMITTED"}
    ca = next(r for r in rows if r["name"] == "Capital Adequacy")
    assert ca["period"] == "Q3 2026" and ca["owner_department"] == "Finance" and ca["frequency"] == "Quarterly"


def test_calendar_seed_is_idempotent(db):
    import psycopg2.extras
    from seed_reports import seed
    cur = db.connection.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    from conftest import TODAY
    seed(cur, today=TODAY)
    cur.execute("select count(*) n from report_instances"); assert cur.fetchone()["n"] == 8
    cur.execute("select count(*) n from report_line_items"); assert cur.fetchone()["n"] == 17
    cur.execute("select count(*) n from calculation_audit"); assert cur.fetchone()["n"] == 17
    cur.execute("select count(*) n from validation_rules"); assert cur.fetchone()["n"] == 4


# ---- the open report ------------------------------------------------------------------------
def test_report_has_the_source_documents_layout(report):
    assert report["available"] is True
    assert [(s["section"], s["title"], len(s["lines"])) for s in report["sections"]] == [
        ("A", "CAPITAL", 9), ("B", "RISK WEIGHTED ASSETS", 4), ("C", "RATIOS", 4)]
    assert [l["line_code"] for s in report["sections"] for l in s["lines"]] == FORM_CODES


def test_report_figures_match_the_source_documents_worked_example(report):
    v = values(report)
    assert (v["A.1"], v["A.2"], v["A.3"], v["A.4"], v["A.5"]) == (150_000_000, 45_000_000, 20_000_000, -8_000_000, 207_000_000)
    assert (v["A.6"], v["A.7"], v["A.8"], v["A.9"]) == (30_000_000, 9_000_000, 39_000_000, 246_000_000)
    assert (v["B.1"], v["B.2"], v["B.3"], v["B.4"]) == (1_650_000_000, 120_000_000, 215_000_000, 1_985_000_000)
    assert percent(v["C.1"]) == "10.43%" and percent(v["C.2"]) == "12.39%" and percent(v["C.3"]) == "12.00%"
    # The source document prints 7,700,000 here; 246,000,000 - 12% x 1,985,000,000 is 7,800,000.
    assert v["C.4"] == 7_800_000


def test_report_uses_the_latest_month_with_a_usable_rwa_not_the_zero_rwa_month(report):
    # capital_positions has 2026-08 with RWA 0 (excluded) and 2026-07: the report is built on 2026-07
    assert values(report)["A.5"] == 207_000_000


def test_only_the_figures_the_schema_lacks_are_flagged_as_demo_inputs(report):
    lines = {l["line_code"]: l for s in report["sections"] for l in s["lines"]}
    demo = {c for c, l in lines.items() if l["is_demo_input"]}
    assert demo == {"A.1", "A.2", "A.3", "A.4", "A.6", "A.7", "B.1", "B.2", "B.3"}
    assert not lines["A.5"]["is_demo_input"] and not lines["B.4"]["is_demo_input"] and not lines["C.2"]["is_demo_input"]


def test_prior_period_comparison_is_honest_when_there_is_no_prior_period(report):
    assert len(report["comparison"]) == 17
    assert all(c["prior"] is None and c["change"] is None and c["change_pct"] is None for c in report["comparison"])
    assert not any(c["needs_explanation"] for c in report["comparison"])


def test_a_report_with_no_figures_says_so_instead_of_erroring(client, auth):
    rows = client.get(f"{API}/reports", headers=auth).json()
    empty = next(r for r in rows if not r["has_lines"])
    body = client.get(f"{API}/reports/{empty['report_instance_id']}", headers=auth).json()
    assert body["available"] is False and body["sections"] == []
    assert client.get(f"{API}/reports/999999", headers=auth).status_code == 404


# ---- validation -----------------------------------------------------------------------------
def test_validation_results_are_green_except_the_thin_capital_buffer_which_is_amber(report):
    results = {r["rule_key"]: r for r in report["validation"]}
    assert [r["level"] for r in report["validation"]] == ["pass", "pass", "pass", "warn"]
    assert "0.39 points above the regulatory minimum" in results["CAPITAL_BUFFER"]["message"]
    assert report["blocked"] is False   # amber needs a comment; only red blocks submission


def test_a_wrong_total_turns_validation_red_and_blocks_submission(client, auth, ca_id, db):
    db.execute("update report_line_items set value = value + 5000 where report_instance_id = %s and line_code = 'A.9'", (ca_id,))
    try:
        body = client.get(f"{API}/reports/{ca_id}", headers=auth).json()
        by_key = {r["rule_key"]: r for r in body["validation"]}
        assert by_key["TIER_TOTAL"]["level"] == "fail" and body["blocked"] is True
        assert "total capital (A.9) shows" in by_key["TIER_TOTAL"]["message"]
        assert by_key["SECTION_TOTALS"]["level"] == "fail"   # C.2 and C.4 no longer agree with A.9 either
    finally:
        db.execute("update report_line_items set value = value - 5000 where report_instance_id = %s and line_code = 'A.9'", (ca_id,))
    assert client.get(f"{API}/reports/{ca_id}", headers=auth).json()["blocked"] is False


def test_a_missing_value_fails_required_fields_and_does_not_crash(report):
    lines = {c: Decimal(1) for c in FORM_CODES}
    lines["B.2"] = None
    rules = [{"rule_key": "REQUIRED_FIELDS", "name": "n", "severity": "BLOCKING"}]
    result = run_validation(rules, lines)[0]
    assert result["level"] == "fail" and "B.2" in result["message"]


def test_zero_rwa_is_reported_not_raised():
    lines = {c: Decimal(1) for c in FORM_CODES}
    lines["B.4"] = Decimal(0)
    rules = [{"rule_key": "SECTION_TOTALS", "name": "n", "severity": "BLOCKING"}]
    assert run_validation(rules, lines)[0]["level"] == "fail"


def test_a_rule_with_no_implemented_check_is_flagged_not_hidden():
    result = run_validation([{"rule_key": "NOPE", "name": "x", "severity": "BLOCKING"}], {})[0]
    assert result["level"] == "warn" and "could not be evaluated" in result["message"]


# ---- drill to source ------------------------------------------------------------------------
def test_every_line_drills_to_a_real_calculation_record(client, auth, ca_id):
    for code in FORM_CODES:
        d = client.get(f"{API}/reports/{ca_id}/drill/{code}", headers=auth).json()
        assert d["formula_text"] and d["source_tables"] and d["calculated_at"], code
        assert d["record_count"] is not None and d["record_count"] >= 0, code
        assert d["label"] and d["value"] is not None, code


def test_a_real_figure_is_traced_to_its_source_row_and_a_demo_input_says_so(client, auth, ca_id):
    a5 = client.get(f"{API}/reports/{ca_id}/drill/A.5", headers=auth).json()
    assert a5["formula_text"] == "capital_positions.tier1_capital" and a5["source_tables"] == ["capital_positions"]
    assert "month 2026-07" in a5["filters_applied"] and a5["record_count"] == 1 and a5["is_demo_input"] is False
    a1 = client.get(f"{API}/reports/{ca_id}/drill/A.1", headers=auth).json()
    assert a1["is_demo_input"] is True and "Demo input" in a1["notes"]


def test_credit_risk_line_shows_the_loan_book_cross_check_and_links_to_the_loans(client, auth, ca_id):
    b1 = client.get(f"{API}/reports/{ca_id}/drill/B.1", headers=auth).json()
    assert b1["record_count"] == 4          # loans behind loan_stage_summary
    assert "Cross-check from the loan book" in b1["notes"] and "Mortgage 35%" in b1["notes"]
    assert "35.00" not in b1["notes"] and "less than 0.01% of this line" in b1["notes"]
    assert "Corporate Loan 100% (default)" in b1["notes"]   # product with no risk_weights row is called out
    assert b1["detail_link"] == "/portfolio"
    assert client.get(f"{API}/reports/{ca_id}/drill/C.2", headers=auth).json()["detail_link"] is None


def test_drill_for_an_unknown_line_or_report_is_404(client, auth, ca_id):
    assert client.get(f"{API}/reports/{ca_id}/drill/Z.9", headers=auth).status_code == 404
    assert client.get(f"{API}/reports/999999/drill/A.1", headers=auth).status_code == 404


# ---- exports: PDF, Excel, and that they agree -----------------------------------------------
@pytest.fixture(scope="module")
def pdf(client, auth, ca_id):
    r = client.post(f"{API}/reports/{ca_id}/export/pdf", headers=auth)
    return r


@pytest.fixture(scope="module")
def xlsx(client, auth, ca_id):
    return client.post(f"{API}/reports/{ca_id}/export/excel", headers=auth)


def pdf_text(response) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)


def test_pdf_is_a_real_pdf_with_the_regulator_layout(pdf):
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    assert pdf.headers["content-disposition"] == 'attachment; filename="Capital_Adequacy_Q3_2026.pdf"'
    assert pdf.content.startswith(b"%PDF")
    text = pdf_text(pdf)
    for expected in ("Bank X", "Capital Adequacy Return", "Q3 2026", "Draft", "SECTION A", "SECTION B", "SECTION C",
                     "TIER 1 CAPITAL", "TOTAL CAPITAL", "TOTAL RWA", "Surplus / (Shortfall)", "Page 1 of 1",
                     "Prepared by", "Reviewed by", "Approved by"):
        assert expected in text, expected
    assert "* Demo input" in text


def test_pdf_shows_every_line_code_and_figure(pdf, report):
    text = pdf_text(pdf)
    for l in (l for s in report["sections"] for l in s["lines"]):
        shown = line_text(l["unit"], l["value"])
        pattern = rf"{re.escape(l['line_code'])}\s+{re.escape(l['label'])}\s+{re.escape(shown)}"
        assert re.search(pattern, text), f"{l['line_code']} {l['label']} {shown}"


def test_pdf_marks_demo_inputs_with_an_asterisk_and_real_figures_without(pdf):
    text = pdf_text(pdf)
    assert re.search(r"A\.1\s+Paid-up capital\s+150,000,000\*", text)
    assert re.search(r"A\.5\s+TIER 1 CAPITAL\s+207,000,000(?!\*)", text)


def test_excel_opens_and_has_the_return_and_its_sources(xlsx):
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert xlsx.headers["content-disposition"] == 'attachment; filename="Capital_Adequacy_Q3_2026.xlsx"'
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Return", "Sources"]
    assert wb["Return"]["A2"].value == "Capital Adequacy Return"
    assert wb["Sources"].max_row == 18            # header + 17 lines
    assert wb["Return"]["C5"].value == "Value"


def test_excel_and_pdf_show_identical_figures_no_drift(pdf, xlsx, report):
    """The spec's acceptance test: every figure reads the same in the screen data, the PDF and the Excel file."""
    ws = load_workbook(io.BytesIO(xlsx.content))["Return"]
    excel = {}
    for row in ws.iter_rows(min_row=6, values_only=True):
        if row[0] and re.fullmatch(r"[ABC]\.\d", str(row[0])):
            excel[row[0]] = row
    text = pdf_text(pdf)
    for l in (l for s in report["sections"] for l in s["lines"]):
        code, unit = l["line_code"], l["unit"]
        from_excel = line_text(unit, Decimal(str(excel[code][2])))
        from_screen = line_text(unit, l["value"])
        assert from_excel == from_screen, code
        assert re.search(rf"{re.escape(code)}\s+{re.escape(l['label'])}\s+{re.escape(from_excel)}", text), code
        assert (excel[code][3] == "Demo input") == l["is_demo_input"], code


def test_excel_numbers_are_real_numbers_with_the_right_formats(xlsx):
    ws = load_workbook(io.BytesIO(xlsx.content))["Return"]
    cells = {ws.cell(r, 1).value: ws.cell(r, 3) for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value}
    assert cells["A.5"].value == 207000000 and isinstance(cells["A.5"].value, (int, float))   # a number, not text
    assert cells["A.4"].value == -8000000 and cells["A.4"].number_format == "#,##0;(#,##0)"
    assert cells["C.2"].number_format == '0.00"%"' and round(cells["C.2"].value, 2) == 12.39


def test_exports_need_a_login_and_fail_cleanly_for_a_report_with_no_figures(client, auth, ca_id):
    assert client.post(f"{API}/reports/{ca_id}/export/pdf").status_code == 401
    assert client.post(f"{API}/reports/{ca_id}/export/excel").status_code == 401
    empty = next(r for r in client.get(f"{API}/reports", headers=auth).json() if not r["has_lines"])
    for kind in ("pdf", "excel"):
        assert client.post(f"{API}/reports/{empty['report_instance_id']}/export/{kind}", headers=auth).status_code == 404


def test_rebuilding_a_report_is_idempotent_and_refreshes_from_the_data(db, ca_id):
    import psycopg2.extras
    from conftest import TODAY
    from app.reports import capital_adequacy
    cur = db.connection.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    db.execute("update capital_positions set tier1_capital = 208000000 where month = '2026-07'")
    try:
        built = capital_adequacy.build(cur, ca_id)
        assert built["A.5"] == 208_000_000 and built["A.9"] == 247_000_000
        assert sum(built[c] for c in ("A.1", "A.2", "A.3", "A.4")) == 208_000_000     # components still add up
        cur.execute("select count(*) n from report_line_items where report_instance_id = %s", (ca_id,)); assert cur.fetchone()["n"] == 17
        cur.execute("select count(*) n from calculation_audit where report_instance_id = %s", (ca_id,)); assert cur.fetchone()["n"] == 17
    finally:
        db.execute("update capital_positions set tier1_capital = 207000000 where month = '2026-07'")
        capital_adequacy.build(cur, ca_id)


def test_building_with_no_usable_capital_data_says_so(db):
    import psycopg2.extras
    from app.reports import capital_adequacy
    cur = db.connection.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    db.execute("create temp table _cp as select * from capital_positions")
    db.execute("delete from capital_positions")
    try:
        with pytest.raises(ValueError, match="no row with a usable risk_weighted_assets"):
            capital_adequacy.build(cur, 1)
    finally:
        db.execute("insert into capital_positions select * from _cp")


# ---- loan drill-down filters (customer, ageing range, LTV bucket) -----------------------------
@pytest.fixture(scope="module")
def ltv_loans(db):
    # collateral: L1 1000 (LTV 90%), L2 400 (112.5%), L3 5000 (38%), L4 none (uncovered); restored afterwards
    db.execute("create temp table _loan_collateral as select loan_id, collateral_value from loans")
    for loan, coll in (("L1", 1000), ("L2", 400), ("L3", 5000), ("L4", None)):
        db.execute("update loans set collateral_value = %s where loan_id = %s", (coll, loan))
    yield
    db.execute("update loans l set collateral_value = t.collateral_value from _loan_collateral t where l.loan_id = t.loan_id")


def ids(client, auth, query):
    return [i["loan_id"] for i in client.get(f"{API}/portfolio/loans?{query}", headers=auth).json()["items"]]


def test_loans_can_be_filtered_by_customer(client, auth):
    assert sorted(ids(client, auth, "customer_id=C1")) == ["L1", "L4"]
    assert ids(client, auth, "customer_id=NOBODY") == []


def test_loans_can_be_filtered_by_a_days_past_due_range_for_the_ageing_buckets(client, auth):
    assert ids(client, auth, "min_days_past_due=0&max_days_past_due=0") == ["L1", "L4"]          # Current
    assert ids(client, auth, "min_days_past_due=30&max_days_past_due=59") == ["L2"]              # 45 days
    assert ids(client, auth, "min_days_past_due=90&max_days_past_due=179") == ["L3"]             # 120 days


@pytest.mark.parametrize("bucket,expected", [("<50%", ["L3"]), ("80-100%", ["L1"]), (">100%", ["L2", "L4"])])
def test_loans_can_be_filtered_by_ltv_bucket_and_missing_collateral_counts_as_uncovered(client, auth, ltv_loans, bucket, expected):
    from urllib.parse import quote
    assert sorted(ids(client, auth, "ltv_bucket=" + quote(bucket))) == expected


def test_ltv_buckets_partition_the_loan_book(client, auth, ltv_loans):
    from urllib.parse import quote
    seen = [i for b in ("<50%", "50-80%", "80-100%", ">100%") for i in ids(client, auth, "ltv_bucket=" + quote(b))]
    assert sorted(seen) == ["L1", "L2", "L3", "L4"]      # every loan lands in exactly one bucket


def test_loan_filters_reject_unknown_ltv_bucket(client, auth):
    assert client.get(f"{API}/portfolio/loans?ltv_bucket=bogus", headers=auth).status_code == 422


# ---- scenarios ------------------------------------------------------------------------------
GOOD = {"name": "  Adverse +5% NPL  ", "inputs": {"devaluation_pct": 20, "rate_change_pct": 2, "npl_increase_pct": 5,
                                                   "deposit_outflow_pct": 10},
        "assumptions": {"coveragePct": 53.5}, "outputs": {"carAfter": 9.1}}


def test_a_scenario_can_be_saved_and_listed_with_who_saved_it(client, auth):
    r = client.post(f"{API}/scenario/save", json=GOOD, headers=auth)
    assert r.status_code == 201 and r.json()["name"] == "Adverse +5% NPL"     # name is trimmed
    saved = client.get(f"{API}/scenario/saved", headers=auth).json()
    row = next(s for s in saved if s["scenario_id"] == r.json()["scenario_id"])
    assert row["inputs"]["devaluation_pct"] == 20 and row["outputs"] == {"carAfter": 9.1}
    assert row["assumptions"] == {"coveragePct": 53.5} and row["created_by_name"] == "Demo Analyst"


@pytest.mark.parametrize("patch", [
    {"name": "   "}, {"name": ""}, {"name": "x" * 81},
    {"inputs": {"devaluation_pct": 51, "rate_change_pct": 0, "npl_increase_pct": 0, "deposit_outflow_pct": 0}},
    {"inputs": {"devaluation_pct": 0, "rate_change_pct": -5.1, "npl_increase_pct": 0, "deposit_outflow_pct": 0}},
    {"inputs": {"devaluation_pct": 0, "rate_change_pct": 0, "npl_increase_pct": 16, "deposit_outflow_pct": 0}},
    {"inputs": {"devaluation_pct": 0, "rate_change_pct": 0, "npl_increase_pct": 0, "deposit_outflow_pct": 31}},
    {"inputs": {"devaluation_pct": 0}},
])
def test_saving_a_scenario_rejects_out_of_range_or_incomplete_input(client, auth, patch):
    assert client.post(f"{API}/scenario/save", json={**GOOD, **patch}, headers=auth).status_code == 422


def test_scenario_endpoints_need_a_login(client):
    assert client.post(f"{API}/scenario/save", json=GOOD).status_code == 401
    assert client.get(f"{API}/scenario/saved").status_code == 401


def test_a_missing_migration_gives_a_clear_503_not_a_500(client, auth, monkeypatch):
    def missing(*_, **__):
        raise psycopg2.errors.UndefinedTable("relation \"saved_scenarios\" does not exist")

    monkeypatch.setattr("app.routers.scenario.query", missing)
    r = client.get(f"{API}/scenario/saved", headers=auth)
    assert r.status_code == 503 and "apply_migration.py" in r.json()["detail"]


# ---- drill-downs and the other exports ------------------------------------------------------
def test_customers_can_be_listed_by_branch_or_segment_with_their_loan_counts(client, auth):
    body = client.get(f"{API}/portfolio/customers?branch_id=B1", headers=auth).json()
    assert body["total"] == 2 and [c["customer_id"] for c in body["items"]] == ["C1", "C3"]   # sorted by name: Ann, Cy
    by_id = {c["customer_id"]: c for c in body["items"]}
    assert by_id["C1"]["loan_count"] == 2 and by_id["C1"]["bad_loan_count"] == 0
    assert by_id["C3"]["loan_count"] == 1 and by_id["C3"]["bad_loan_count"] == 1        # L3 is 120 days past due
    assert [c["customer_id"] for c in client.get(f"{API}/portfolio/customers?segment=SME", headers=auth).json()["items"]] == ["C2"]
    assert client.get(f"{API}/portfolio/customers?branch_id=NOPE", headers=auth).json()["total"] == 0
    assert client.get(f"{API}/portfolio/customers?limit=0", headers=auth).status_code == 422


def test_channel_usage_is_a_current_snapshot_of_transaction_counts(client, auth):
    body = client.get(f"{API}/performance/channels", headers=auth).json()
    counts = {c["channel"]: c["transaction_count"] for c in body["channels"]}
    assert counts == {"ATM": 3, "Branch": 2, "Mobile": 1}
    assert [c["channel"] for c in body["channels"]] == ["ATM", "Branch", "Mobile"]        # biggest first
    assert abs(sum(c["share_pct"] for c in body["channels"]) - 100) < 0.2
    assert body["first_date"] == "2026-09-20" and body["last_date"] == "2026-09-21"


def test_portfolio_excel_has_the_sheets_and_the_same_numbers_as_the_screen(client, auth):
    r = client.get(f"{API}/portfolio/export.xlsx", headers=auth)
    assert r.status_code == 200 and "portfolio_credit_risk.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["IFRS 9 stages", "By product", "By segment", "By branch", "By currency",
                             "Top exposures", "Ageing", "LTV"]
    stages = [row for row in wb["IFRS 9 stages"].iter_rows(min_row=2, values_only=True)]
    assert [s[0] for s in stages] == [1, 2, 3] and stages[2][2] == 1900.0
    assert [row[0] for row in wb["Ageing"].iter_rows(min_row=2, values_only=True)] == ["Current", "31-60", "90-180", "180+"]
    assert [row[0] for row in wb["By product"].iter_rows(min_row=2, values_only=True)] == ["Corporate Loan", "Mortgage"]


def test_performance_excel_includes_the_regional_rollup_and_the_allocation_label(client, auth):
    r = client.get(f"{API}/performance/export.xlsx", headers=auth)
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Branches", "Regions", "Segments", "Products", "Channels"]
    regions = {row[0]: row for row in wb["Regions"].iter_rows(min_row=2, values_only=True)}
    assert set(regions) == {"Lebanon", "KSA"} and regions["Lebanon"][4] == 50.0
    assert "allocated proportionally" in wb["Segments"]["G1"].value
    assert [row[0] for row in wb["Branches"].iter_rows(min_row=2, values_only=True)] == ["B1", "B2"]


def test_new_endpoints_need_a_login(client):
    for path in ("/reports", "/reports/1", "/reports/1/drill/A.1", "/portfolio/customers", "/portfolio/export.xlsx",
                 "/performance/channels", "/performance/export.xlsx"):
        assert client.get(API + path).status_code == 401, path
